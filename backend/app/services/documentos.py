"""
Proyectos de dictamen editables dentro del sistema.

El proyecto que viaja "a la firma" se guarda como HTML (lo que edita el editor
del navegador) + un Word "base" (la plantilla ya completada, o el Word que se
adjuntó). La idea central es NO reconstruir el Word desde cero:

  - Cada párrafo del HTML recuerda de qué párrafo del Word vino (data-p="n").
    Al armar el Word, si el párrafo no se tocó se copia TAL CUAL del original
    (letra, tamaño, espaciado, sangrías, numeración, logos flotantes...). Si se
    corrigió, se rearma con el mismo formato de párrafo y de letra del original.
  - Las imágenes (logos, sellos) se muestran en el editor pero viajan como
    referencia (data-img="párrafo:run"): en el Word se copia la imagen original.
  - Lo que el editor no sabe editar (tablas con imágenes, controles de
    contenido) viaja como "bloque fijo" (data-b="n") y se copia intacto.
  - El encabezado/pie (membrete) queda siempre el del Word base.
  - El PDF se arma convirtiendo ese Word con LibreOffice (si está instalado,
    como en el servidor); si no, con reportlab (aproximado).

  docx_a_html(bytes)            → HTML para el editor
  membrete_html(bytes)          → {"encabezado", "pie"} para mostrar arriba/abajo
  reemplazar_variables(bytes, valores) → Word con @numero, @caratula... completados
  sanitizar(html)               → HTML limpio (solo lo permitido)
  armar_docx(base, html)        → bytes de un .docx
  armar_pdf(base, html)         → bytes de un .pdf
  texto_plano / formato_plano   → para comparar versiones

El editor trabaja con `white-space: pre-wrap`: los espacios y tabulaciones del
texto son reales (no se colapsan como en HTML común).
"""

import copy
import hashlib
import io
import os
import re
import shutil
import subprocess
import tempfile
import threading
from collections import Counter
from html import escape
from html.parser import HTMLParser

# ═══════════════════════════════════════════════════════════════
# HTML → árbol simple
# ═══════════════════════════════════════════════════════════════

_VACIAS = {"br", "img", "hr", "meta", "link", "input", "col"}
_DESCARTAR_CONTENIDO = {"script", "style", "title", "head", "xml", "template", "noscript", "iframe", "object"}


class _Nodo:
    __slots__ = ("tag", "attrs", "hijos", "texto")

    def __init__(self, tag=None, attrs=None, texto=None):
        self.tag = tag
        self.attrs = attrs or {}
        self.hijos = []
        self.texto = texto


class _Parser(HTMLParser):
    def __init__(self):
        super().__init__(convert_charrefs=True)
        self.raiz = _Nodo("root")
        self.pila = [self.raiz]
        self.saltar = 0

    def handle_starttag(self, tag, attrs):
        tag = tag.lower()
        if ":" in tag:  # <o:p>, <v:shape> de Word
            return
        if tag in _DESCARTAR_CONTENIDO:
            self.saltar += 1
            return
        if self.saltar:
            return
        n = _Nodo(tag, {k.lower(): (v or "") for k, v in attrs})
        self.pila[-1].hijos.append(n)
        if tag not in _VACIAS:
            self.pila.append(n)

    def handle_startendtag(self, tag, attrs):
        tag = tag.lower()
        if tag in _DESCARTAR_CONTENIDO or self.saltar or ":" in tag:
            return
        self.pila[-1].hijos.append(_Nodo(tag, {k.lower(): (v or "") for k, v in attrs}))

    def handle_endtag(self, tag):
        tag = tag.lower()
        if ":" in tag:
            return
        if tag in _DESCARTAR_CONTENIDO:
            self.saltar = max(0, self.saltar - 1)
            return
        if self.saltar or tag in _VACIAS:
            return
        for i in range(len(self.pila) - 1, 0, -1):
            if self.pila[i].tag == tag:
                del self.pila[i:]
                break

    def handle_data(self, data):
        if self.saltar or not data:
            return
        self.pila[-1].hijos.append(_Nodo(texto=data))


def _parsear(html: str) -> _Nodo:
    p = _Parser()
    p.feed(html or "")
    p.close()
    return p.raiz


def _estilo(attrs) -> dict:
    out = {}
    for parte in (attrs.get("style") or "").split(";"):
        if ":" in parte:
            k, v = parte.split(":", 1)
            out[k.strip().lower()] = v.strip().lower()
    return out


def _alineacion(attrs) -> str | None:
    a = (_estilo(attrs).get("text-align") or attrs.get("align") or "").lower()
    if a in ("center", "right", "justify", "left"):
        return a
    if a == "start":
        return "left"
    if a == "end":
        return "right"
    return None


def _medida_pt(valor: str | None, solo_pt: bool = False) -> float | None:
    """'36pt' / '48px' / '1.27cm' / '2em' → puntos."""
    if not valor:
        return None
    m = re.match(r"^\s*(-?[\d.]+)\s*(pt|px|cm|mm|in|em|rem)?\s*$", valor)
    if not m:
        return None
    u = m.group(2) or "px"
    if solo_pt and u != "pt":
        return None
    try:
        n = float(m.group(1))
    except ValueError:
        return None
    return round({"pt": n, "px": n * 0.75, "cm": n * 28.3465, "mm": n * 2.83465,
                  "in": n * 72, "em": n * 12, "rem": n * 12}[u], 1)


def _interlineado_css(valor: str | None):
    """line-height del HTML → múltiplo (float) o ("pt", n)."""
    if not valor:
        return None
    v = valor.strip()
    if v.endswith("pt"):
        n = _medida_pt(v, solo_pt=True)
        return ("pt", n) if n else None
    try:
        return round(float(v) / 1.15, 2)
    except ValueError:
        return None


def _entero(v) -> int | None:
    return int(v) if isinstance(v, str) and v.isdigit() else None


# ═══════════════════════════════════════════════════════════════
# Sanitizar: deja solo lo que el editor y los exportadores entienden
# ═══════════════════════════════════════════════════════════════

_INLINE_OK = {"b": "b", "strong": "b", "i": "i", "em": "i", "u": "u", "ins": "u",
              "s": "s", "strike": "s", "del": "s", "sup": "sup", "sub": "sub"}
_BLOQUE_OK = {"p": "p", "div": "p", "section": "p", "article": "p", "h1": "h2", "h2": "h2",
              "h3": "h3", "h4": "h3", "h5": "h3", "h6": "h3", "li": "li", "blockquote": "blockquote",
              "ul": "ul", "ol": "ol", "table": "table", "thead": "tbody", "tbody": "tbody",
              "tfoot": "tbody", "tr": "tr", "td": "td", "th": "td"}
_CONTENEDORES = {"root", "html", "body", "ul", "ol", "table", "thead", "tbody", "tfoot", "tr", "blockquote"}
_SRC_OK = re.compile(r"^/uploads/[0-9a-f]{32}\.(png|jpg|gif)$")
_IMG_OK = re.compile(r"^\d+:\d+$")


def _estilo_bloque(attrs) -> str:
    partes = []
    al = _alineacion(attrs)
    if al and al != "left":
        partes.append(f"text-align:{al}")
    est = _estilo(attrs)
    ml = _medida_pt(est.get("margin-left") or est.get("padding-left"))
    if ml and ml > 0:
        partes.append(f"margin-left:{ml}pt")
    ti = _medida_pt(est.get("text-indent"))
    if ti:
        partes.append(f"text-indent:{ti}pt")
    # Espaciado que viene del Word (solo en pt: lo que pone el navegador en px se ignora)
    for k in ("margin-top", "margin-bottom"):
        v = _medida_pt(est.get(k), solo_pt=True)
        if v:
            partes.append(f"{k}:{v}pt")
    lh = est.get("line-height", "")
    if re.match(r"^\d+(\.\d+)?(pt)?$", lh):
        partes.append(f"line-height:{lh}")
    return ";".join(partes)


def _estilo_img(attrs) -> str:
    est = _estilo(attrs)
    partes = []
    for k in ("width", "height", "margin-left"):
        v = _medida_pt(est.get(k), solo_pt=True)
        if v:
            partes.append(f"{k}:{v}pt")
    if est.get("float") in ("left", "right"):
        partes.append(f"float:{est['float']}")
    if est.get("display") == "block":
        partes.append("display:block")
    return ";".join(partes)


def _serializar(nodo: _Nodo, out: list):
    for h in nodo.hijos:
        if h.tag is None:
            # Saltos de línea sueltos entre bloques: con pre-wrap se verían como renglones vacíos
            if nodo.tag in _CONTENEDORES and not h.texto.strip():
                continue
            out.append(escape(h.texto, quote=False))
            continue
        t = h.tag
        if t == "br":
            out.append("<br>")
        elif t == "img":
            clave = h.attrs.get("data-img", "")
            src = h.attrs.get("data-src", "")
            if not _IMG_OK.match(clave) and not _SRC_OK.match(src):
                continue  # imágenes pegadas de afuera: no se aceptan
            extra = f' data-img="{clave}"' if _IMG_OK.match(clave) else ""
            if _SRC_OK.match(src):
                extra += f' data-src="{src}"'
            est = _estilo_img(h.attrs)
            if est:
                extra += f' style="{est}"'
            alt = h.attrs.get("alt", "")
            if alt:
                extra += f' alt="{escape(alt[:40])}"'
            out.append(f'<img{extra} contenteditable="false">')
        elif t in _INLINE_OK:
            tag = _INLINE_OK[t]
            out.append(f"<{tag}>")
            _serializar(h, out)
            out.append(f"</{tag}>")
        elif t == "span":
            est = _estilo(h.attrs)
            envolver = []
            if est.get("font-weight") in ("bold", "700", "800", "900", "bolder"):
                envolver.append("b")
            if "italic" in est.get("font-style", ""):
                envolver.append("i")
            if "underline" in est.get("text-decoration", "") or "underline" in est.get("text-decoration-line", ""):
                envolver.append("u")
            tam = _medida_pt(est.get("font-size"), solo_pt=True)
            if tam:
                out.append(f'<span style="font-size:{tam}pt">')
            out.extend(f"<{e}>" for e in envolver)
            _serializar(h, out)
            out.extend(f"</{e}>" for e in reversed(envolver))
            if tam:
                out.append("</span>")
        elif t == "div" and _entero(h.attrs.get("data-b")) is not None:
            out.append(f'<div data-b="{int(h.attrs["data-b"])}" contenteditable="false" class="bloque-fijo">')
            _serializar(h, out)
            out.append("</div>")
        elif t in _BLOQUE_OK:
            tag = _BLOQUE_OK[t]
            extra = ""
            if tag in ("p", "h2", "h3", "li") and _entero(h.attrs.get("data-p")) is not None:
                extra += f' data-p="{int(h.attrs["data-p"])}"'
            if tag == "table" and _entero(h.attrs.get("data-t")) is not None:
                extra += f' data-t="{int(h.attrs["data-t"])}"'
            if tag in ("p", "h2", "h3", "li", "blockquote", "td"):
                est = _estilo_bloque(h.attrs)
                if est:
                    extra += f' style="{est}"'
            if tag == "td":
                cs = h.attrs.get("colspan", "")
                if cs.isdigit() and int(cs) > 1:
                    extra += f' colspan="{int(cs)}"'
            out.append(f"<{tag}{extra}>")
            _serializar(h, out)
            out.append(f"</{tag}>")
        else:
            # Etiqueta desconocida (a, font...): se queda solo el contenido
            _serializar(h, out)


def sanitizar(html: str) -> str:
    out = []
    _serializar(_parsear(html), out)
    return "".join(out)


# ═══════════════════════════════════════════════════════════════
# HTML → bloques (estructura intermedia para Word y PDF)
# ═══════════════════════════════════════════════════════════════
# Párrafo: {"tipo": "p", "origen", "align", "izq", "primera", "antes", "despues", "interl",
#           "titulo": 0|1|2, "lista": None|"ul"|"ol", "numero",
#           "runs": [{"t", "b", "i", "u", "s", "sup", "sub", "sz"} | {"br": True} | {"img": "p:r", "src", "w", "h"}]}
# Tabla:   {"tipo": "tabla", "origen", "filas": [[{"bloques": [...], "colspan": n}]]}
# Fijo:    {"tipo": "fijo", "origen", "bloques": [...]}  (se copia intacto del Word)

class _Armador:
    def __init__(self):
        self.bloques = []
        self.actual = None

    def cerrar(self, forzar=False):
        a = self.actual
        if a is not None:
            while a["runs"] and a["runs"][-1].get("br"):
                a["runs"].pop()
                forzar = True
            if a["runs"] or forzar or a.get("explicito"):
                a.pop("explicito", None)
                self.bloques.append(a)
        self.actual = None

    def abrir(self, props):
        self.cerrar()
        self.actual = {"tipo": "p", "runs": [], **props}

    def _asegurar(self, props):
        if self.actual is None:
            self.actual = {"tipo": "p", "runs": [], **props}

    def texto(self, t, fmt, props):
        if self.actual is None and not t.strip():
            return  # espacios sueltos entre bloques
        self._asegurar(props)
        partes = t.replace("\r\n", "\n").replace("\r", "\n").split("\n")
        for i, parte in enumerate(partes):
            if i:
                self.actual["runs"].append({"br": True})
            if parte:
                self.actual["runs"].append({"t": parte.replace("\xa0", " "), **fmt})

    def salto(self, props):
        self._asegurar(props)
        self.actual["runs"].append({"br": True})

    def imagen(self, run, props):
        self._asegurar(props)
        self.actual["runs"].append(run)


def _props_de(attrs, base) -> dict:
    p = dict(base)
    al = _alineacion(attrs)
    if al:
        p["align"] = al
    est = _estilo(attrs)
    ml = _medida_pt(est.get("margin-left"))
    if ml:
        p["izq"] = (base.get("izq") or 0) + ml
    ti = _medida_pt(est.get("text-indent"))
    if ti is not None:
        p["primera"] = ti
    for k, clave in (("margin-top", "antes"), ("margin-bottom", "despues")):
        v = _medida_pt(est.get(k), solo_pt=True)
        if v is not None:
            p[clave] = v
    il = _interlineado_css(est.get("line-height"))
    if il is not None:
        p["interl"] = il
    return p


_INLINE_FMT = {"b": "b", "strong": "b", "i": "i", "em": "i", "u": "u", "ins": "u",
               "s": "s", "strike": "s", "del": "s", "sup": "sup", "sub": "sub"}


def _recorrer(nodo, arm: _Armador, fmt: dict, props: dict, ctx_lista=None):
    for h in nodo.hijos:
        if h.tag is None:
            arm.texto(h.texto, fmt, props)
            continue
        t = h.tag
        if t == "br":
            arm.salto(props)
        elif t == "img":
            clave = h.attrs.get("data-img", "")
            if _IMG_OK.match(clave):
                est = _estilo(h.attrs)
                arm.imagen({"img": clave, "src": h.attrs.get("data-src") or None,
                            "w": _medida_pt(est.get("width")), "h": _medida_pt(est.get("height"))}, props)
        elif t in _INLINE_FMT or t == "span":
            f = dict(fmt)
            if t in _INLINE_FMT:
                f[_INLINE_FMT[t]] = True
            else:
                est = _estilo(h.attrs)
                if est.get("font-weight") in ("bold", "700", "800", "900", "bolder"):
                    f["b"] = True
                if "italic" in est.get("font-style", ""):
                    f["i"] = True
                if "underline" in est.get("text-decoration", ""):
                    f["u"] = True
                tam = _medida_pt(est.get("font-size"), solo_pt=True)
                if tam:
                    f["sz"] = tam
            _recorrer(h, arm, f, props, ctx_lista)
        elif t == "div" and _entero(h.attrs.get("data-b")) is not None:
            arm.cerrar()
            sub = _Armador()
            _recorrer(h, sub, {}, {}, None)
            sub.cerrar()
            arm.bloques.append({"tipo": "fijo", "origen": int(h.attrs["data-b"]), "bloques": sub.bloques})
        elif t in ("p", "div", "section", "article", "h1", "h2", "h3", "h4", "h5", "h6"):
            pp = _props_de(h.attrs, {k: v for k, v in props.items() if k in ("izq", "align")})
            if t in ("h1", "h2"):
                pp["titulo"] = 1
            elif t in ("h3", "h4", "h5", "h6"):
                pp["titulo"] = 2
            if _entero(h.attrs.get("data-p")) is not None:
                pp["origen"] = int(h.attrs["data-p"])
            pp["explicito"] = True
            arm.abrir(pp)
            _recorrer(h, arm, fmt, pp, ctx_lista)
            arm.cerrar()
        elif t == "blockquote":
            arm.cerrar()
            pp = _props_de(h.attrs, props)
            pp["izq"] = (props.get("izq") or 0) + 36
            pp.pop("origen", None)
            _recorrer(h, arm, fmt, pp, ctx_lista)
            arm.cerrar()
        elif t in ("ul", "ol"):
            arm.cerrar()
            lista = {"tipo": t, "n": 0}
            pp = {k: v for k, v in props.items() if k != "origen"}
            if ctx_lista:  # lista anidada
                pp["izq"] = (props.get("izq") or 0) + 18
            _recorrer(h, arm, fmt, pp, lista)
            arm.cerrar()
        elif t == "li":
            pp = _props_de(h.attrs, props)
            if ctx_lista:
                ctx_lista["n"] += 1
                pp["lista"] = ctx_lista["tipo"]
                pp["numero"] = ctx_lista["n"]
            if _entero(h.attrs.get("data-p")) is not None:
                pp["origen"] = int(h.attrs["data-p"])
            pp["explicito"] = True
            arm.abrir(pp)
            _recorrer(h, arm, fmt, pp, None)
            arm.cerrar()
        elif t == "table":
            arm.cerrar()
            filas = []
            for tr in _buscar(h, "tr"):
                fila = []
                for td in tr.hijos:
                    if td.tag in ("td", "th"):
                        sub = _Armador()
                        _recorrer(td, sub, dict(fmt, b=True) if td.tag == "th" else fmt,
                                  _props_de(td.attrs, {}), None)
                        sub.cerrar()
                        cs = td.attrs.get("colspan", "")
                        fila.append({"bloques": sub.bloques, "colspan": int(cs) if cs.isdigit() else 1})
                if fila:
                    filas.append(fila)
            if filas:
                arm.bloques.append({"tipo": "tabla", "filas": filas, "origen": _entero(h.attrs.get("data-t"))})
        elif t == "hr":
            continue
        else:
            _recorrer(h, arm, fmt, props, ctx_lista)


def _buscar(nodo, tag):
    for h in nodo.hijos:
        if h.tag == tag:
            yield h
        elif h.tag in ("thead", "tbody", "tfoot"):
            yield from _buscar(h, tag)


def html_a_bloques(html: str) -> list:
    arm = _Armador()
    _recorrer(_parsear(html), arm, {}, {})
    arm.cerrar()
    return arm.bloques


def _firma(b):
    """Lo editable de un bloque: si coincide con el original, el párrafo no se tocó."""
    if b["tipo"] == "tabla":
        return ("t", tuple(tuple((tuple(_firma(x) for x in c["bloques"]), c["colspan"]) for c in fila)
                           for fila in b["filas"]))
    if b["tipo"] == "fijo":
        return ("f", b.get("origen"))
    runs = []
    for r in b["runs"]:
        if "img" in r:
            runs.append(("img", r["img"]))
        elif r.get("br"):
            runs.append(("br",))
        else:
            k = (bool(r.get("b")), bool(r.get("i")), bool(r.get("u")), bool(r.get("s")),
                 bool(r.get("sup")), bool(r.get("sub")), r.get("sz"))
            if runs and runs[-1][0] == "t" and runs[-1][1] == k:
                runs[-1] = ("t", k, runs[-1][2] + r["t"])
            else:
                runs.append(("t", k, r["t"]))
    al = b.get("align")
    return ("p", None if al in (None, "left") else al, round(b.get("izq") or 0), round(b.get("primera") or 0),
            b.get("lista"), b.get("titulo") or 0, tuple(runs))


def texto_plano(html: str) -> str:
    """Texto del documento (para saber si hubo cambios entre versiones)."""
    lineas = []

    def txt(b):
        if b["tipo"] == "tabla":
            for fila in b["filas"]:
                for c in fila:
                    for x in c["bloques"]:
                        txt(x)
        elif b["tipo"] == "fijo":
            for x in b["bloques"]:
                txt(x)
        else:
            lineas.append("".join(r.get("t", "\n" if r.get("br") else "") for r in b["runs"]))

    for b in html_a_bloques(html):
        txt(b)
    return re.sub(r"[ \t\xa0]+", " ", "\n".join(lineas)).strip()


def formato_plano(html: str) -> str:
    """Texto + formato: cambia si cambió una negrita o una alineación."""
    return sanitizar(html or "").replace("\n", "")


# ═══════════════════════════════════════════════════════════════
# LibreOffice (en el servidor): PDF fiel y vista previa de imágenes EMF/WMF
# ═══════════════════════════════════════════════════════════════

_LO_LOCK = threading.Lock()


def _soffice():
    return shutil.which("soffice") or shutil.which("libreoffice")


def convertir_libreoffice(datos: bytes, ext_entrada: str, formato: str, timeout: int = 90) -> bytes | None:
    """Convierte con LibreOffice sin pantalla. None si no está o falla."""
    exe = _soffice()
    if not exe:
        return None
    with _LO_LOCK, tempfile.TemporaryDirectory() as d:
        entrada = os.path.join(d, "entrada" + ext_entrada)
        with open(entrada, "wb") as f:
            f.write(datos)
        perfil = "file:///" + os.path.join(tempfile.gettempdir(), "lo_perfil_defensoria").replace("\\", "/").lstrip("/")
        try:
            subprocess.run(
                [exe, "--headless", "--norestore", "--nologo", "--nodefault", "--nolockcheck",
                 f"-env:UserInstallation={perfil}", "--convert-to", formato, "--outdir", d, entrada],
                timeout=timeout, capture_output=True,
            )
        except Exception as e:
            print(f"[!] LibreOffice no pudo convertir: {e}")
            return None
        salida = os.path.join(d, "entrada." + formato.split(":")[0])
        if os.path.isfile(salida):
            with open(salida, "rb") as f:
                return f.read()
    return None


# ═══════════════════════════════════════════════════════════════
# Word → HTML
# ═══════════════════════════════════════════════════════════════

def _qn(t):
    from docx.oxml.ns import qn
    return qn(t)


_NS_A = "http://schemas.openxmlformats.org/drawingml/2006/main"
_NS_R = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
_NS_WP = "http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing"
_NS_V = "urn:schemas-microsoft-com:vml"
_NS_MC = "http://schemas.openxmlformats.org/markup-compatibility/2006"
_CONT_RUNS = ("hyperlink", "ins", "smartTag", "fldSimple", "customXml", "dir", "bdo", "sdt", "sdtContent")


def _iter_runs(p_el):
    """Runs de un párrafo en orden (incluye hipervínculos y revisiones insertadas;
    no entra en los cuadros de texto ni en lo borrado)."""
    conts = {_qn("w:" + c) for c in _CONT_RUNS}
    r_tag = _qn("w:r")
    for h in p_el.iterchildren():
        if h.tag == r_tag:
            yield h
        elif h.tag in conts:
            yield from _iter_runs(h)


def _es_especial(r_el) -> bool:
    """Run con dibujo, imagen, objeto o cuadro de texto."""
    for h in r_el.iterchildren():
        if h.tag in (_qn("w:drawing"), _qn("w:pict"), _qn("w:object"), f"{{{_NS_MC}}}AlternateContent"):
            return True
    return False


_CACHE_PREVIEW = {}


def _url_preview(blob: bytes, nombre_parte: str) -> str | None:
    """Sube una copia visible (PNG/JPG) de la imagen y devuelve /uploads/<hash>.ext."""
    from app.services import storage
    clave = hashlib.sha1(blob).hexdigest()[:32]
    if clave in _CACHE_PREVIEW:
        return _CACHE_PREVIEW[clave]
    ext = os.path.splitext(nombre_parte)[1].lower()
    datos, ext_out = None, None
    if ext in (".png", ".gif", ".jpg", ".jpeg"):
        datos, ext_out = blob, (".jpg" if ext == ".jpeg" else ext)
    else:
        try:
            from PIL import Image
            img = Image.open(io.BytesIO(blob))  # EMF/WMF solo se puede en Windows
            buf = io.BytesIO()
            img.save(buf, "PNG")
            datos, ext_out = buf.getvalue(), ".png"
        except Exception:
            if ext in (".emf", ".wmf"):
                datos = convertir_libreoffice(blob, ext, "png", timeout=40)
                ext_out = ".png"
    if not datos:
        _CACHE_PREVIEW[clave] = None
        return None
    nombre = f"{clave}{ext_out}"
    try:
        storage.guardar(nombre, datos, {"png": "image/png", "jpg": "image/jpeg", "gif": "image/gif"}[ext_out[1:]])
    except Exception as e:
        print(f"[!] No se pudo guardar la vista previa de una imagen: {e}")
        return None
    url = f"/uploads/{nombre}"
    _CACHE_PREVIEW[clave] = url
    return url


def _img_html(r_el, part, clave: str, previews: bool) -> str:
    """<img> que representa un run con dibujo (la imagen real queda en el Word)."""
    blob, nombre = None, ""
    for blip in r_el.iter(f"{{{_NS_A}}}blip"):
        rid = blip.get(f"{{{_NS_R}}}embed")
        if rid and rid in part.related_parts:
            p = part.related_parts[rid]
            blob, nombre = p.blob, str(p.partname)
            break
    if blob is None:
        for vi in r_el.iter(f"{{{_NS_V}}}imagedata"):
            rid = vi.get(f"{{{_NS_R}}}id")
            if rid and rid in part.related_parts:
                p = part.related_parts[rid]
                blob, nombre = p.blob, str(p.partname)
                break
    estilo = []
    ext = next(r_el.iter(f"{{{_NS_WP}}}extent"), None)
    if ext is not None:
        try:
            estilo.append(f"width:{round(int(ext.get('cx')) / 12700, 1)}pt")
            estilo.append(f"height:{round(int(ext.get('cy')) / 12700, 1)}pt")
        except (TypeError, ValueError):
            pass
    ancla = next(r_el.iter(f"{{{_NS_WP}}}anchor"), None)
    if ancla is not None:
        # Imagen flotante (ej. el sello al lado de la fecha): se muestra aparte
        pos = ancla.find(f"{{{_NS_WP}}}positionH")
        alinear = pos.find(f"{{{_NS_WP}}}align") if pos is not None else None
        off = pos.find(f"{{{_NS_WP}}}posOffset") if pos is not None else None
        estilo.append("display:block")
        if alinear is not None and alinear.text == "right":
            estilo.append("float:right")
        elif off is not None and (off.text or "").lstrip("-").isdigit() and int(off.text) > 0:
            estilo.append(f"margin-left:{round(int(off.text) / 12700, 1)}pt")
    src = ""
    if blob is not None and previews:
        url = _url_preview(blob, nombre)
        if url:
            src = f' data-src="{url}"'
    alt = "imagen" if blob is not None else "forma"
    st = f' style="{";".join(estilo)}"' if estilo else ""
    dimg = f' data-img="{clave}"' if clave else ""
    return f'<img{dimg}{src}{st} alt="{alt}" contenteditable="false">'


class _Padre:
    """Padre mínimo para los objetos de python-docx (necesitan .part para los estilos)."""
    def __init__(self, part):
        self.part = part


def _cadena_estilos(par, run=None):
    estilos = []
    for est in ((run.style if run is not None else None), par.style):
        s = est
        while s is not None:
            estilos.append(s)
            s = s.base_style
    return estilos


def _run_efectivo(run, par, atributo):
    """Negrita/cursiva/subrayado: directo en el run, o heredado del estilo."""
    v = getattr(run.font, atributo)
    if v is not None:
        return bool(v)
    for s in _cadena_estilos(par, run):
        try:
            sv = getattr(s.font, atributo)
        except Exception:
            sv = None
        if sv is not None:
            return bool(sv)
    return False


def _tam_efectivo(run, par, tam_defecto):
    if run.font.size is not None:
        return run.font.size.pt
    for s in _cadena_estilos(par, run):
        try:
            if s.font.size is not None:
                return s.font.size.pt
        except Exception:
            continue
    return tam_defecto


def _heredado(par, attr):
    v = getattr(par.paragraph_format, attr)
    s = par.style
    while v is None and s is not None:
        try:
            v = getattr(s.paragraph_format, attr)
        except Exception:
            v = None
        s = s.base_style
    return v


def _alineacion_par(par) -> str | None:
    a = _heredado(par, "alignment")
    if a is None:
        return None
    n = str(a)
    if "CENTER" in n:
        return "center"
    if "RIGHT" in n:
        return "right"
    if "JUSTIFY" in n or "DISTRIBUTE" in n:
        return "justify"
    return None


def _indent_par(par, attr):
    v = _heredado(par, attr)
    return round(v.pt, 1) if v is not None else None


def _defaults(doc) -> dict:
    """Valores por defecto del documento (docDefaults)."""
    d = getattr(doc, "_defensoria_defaults", None)
    if d is not None:
        return d
    d = {"tam": 12.0, "antes": 0.0, "despues": 0.0, "interl": 1.0}
    try:
        st = doc.styles.element
        sz = st.xpath("./w:docDefaults/w:rPrDefault/w:rPr/w:sz/@w:val")
        if sz:
            d["tam"] = int(sz[0]) / 2
        else:
            d["tam"] = 10.0  # lo que asume Word si no se indica
        sp = st.xpath("./w:docDefaults/w:pPrDefault/w:pPr/w:spacing")
        if sp:
            s = sp[0]
            if s.get(_qn("w:after")) is not None:
                d["despues"] = int(s.get(_qn("w:after"))) / 20
            if s.get(_qn("w:before")) is not None:
                d["antes"] = int(s.get(_qn("w:before"))) / 20
            if s.get(_qn("w:line")) is not None:
                regla = s.get(_qn("w:lineRule")) or "auto"
                n = int(s.get(_qn("w:line")))
                d["interl"] = round(n / 240, 2) if regla == "auto" else ("pt", n / 20)
    except Exception:
        pass
    try:
        normal = doc.styles["Normal"]
        if normal.font.size is not None:
            d["tam"] = normal.font.size.pt
    except Exception:
        pass
    doc._defensoria_defaults = d
    return d


def _espaciado(par, doc):
    """(antes, después, interlineado) efectivos del párrafo."""
    dfl = _defaults(doc)
    antes = _heredado(par, "space_before")
    despues = _heredado(par, "space_after")
    ls = _heredado(par, "line_spacing")
    rule = _heredado(par, "line_spacing_rule")
    antes = antes.pt if antes is not None else dfl["antes"]
    despues = despues.pt if despues is not None else dfl["despues"]
    if ls is None:
        interl = dfl["interl"]
    elif isinstance(ls, float):
        interl = round(ls, 2)
    else:
        interl = ("pt", round(ls.pt, 1)) if "EXACTLY" in str(rule) or "AT_LEAST" in str(rule) else 1.0
    return antes, despues, interl


def _tipo_lista(par, doc) -> str | None:
    ppr = par._p.pPr
    num_pr = ppr.numPr if ppr is not None else None
    if num_pr is None or num_pr.numId is None:
        # La numeración puede venir del estilo (ej. "List Bullet")
        s = par.style
        num_pr = None
        while s is not None and num_pr is None:
            spr = s.element.pPr
            if spr is not None and spr.numPr is not None and spr.numPr.numId is not None:
                num_pr = spr.numPr
            s = s.base_style
        if num_pr is None:
            return None
    try:
        num_id = num_pr.numId.val
        if num_id == 0:
            return None
        ilvl = num_pr.ilvl.val if num_pr.ilvl is not None else 0
        numbering = doc.part.numbering_part.element
        num = numbering.xpath(f'./w:num[@w:numId="{num_id}"]')
        if not num:
            return None
        abs_id = num[0].xpath("./w:abstractNumId/@w:val")[0]
        fmt = numbering.xpath(
            f'./w:abstractNum[@w:abstractNumId="{abs_id}"]/w:lvl[@w:ilvl="{ilvl}"]/w:numFmt/@w:val')
        if fmt and fmt[0] == "none":
            return None
        return "ul" if (fmt and fmt[0] == "bullet") else "ol"
    except Exception:
        return "ul"


def _par_a_html(p_el, doc, part, idx=None, previews=True) -> tuple[str | None, str]:
    """(tipo_lista, html) de un párrafo de Word. idx: posición en el cuerpo (o None)."""
    from docx.text.paragraph import Paragraph
    from docx.text.run import Run
    par = Paragraph(p_el, _Padre(part))
    tam_def = _defaults(doc)["tam"]
    # Word suele partir una palabra en varios runs con el mismo formato: se unen
    segmentos = []  # (clave_formato, texto) | ("img", html)
    for j, r_el in enumerate(_iter_runs(p_el)):
        if _es_especial(r_el):
            segmentos.append(("img", _img_html(r_el, part, f"{idx}:{j}" if idx is not None else "", previews)))
            continue
        run = Run(r_el, par)
        t = run.text
        if not t:
            continue
        tam = _tam_efectivo(run, par, tam_def)
        clave = (_run_efectivo(run, par, "bold"), _run_efectivo(run, par, "italic"),
                 _run_efectivo(run, par, "underline"), bool(run.font.strike),
                 bool(run.font.superscript), bool(run.font.subscript) and not run.font.superscript,
                 round(tam, 1) if tam and abs(tam - 12) > 0.01 else None)
        if segmentos and segmentos[-1][0] == clave:
            segmentos[-1] = (clave, segmentos[-1][1] + t)
        else:
            segmentos.append((clave, t))
    partes = []
    for clave, t in segmentos:
        if clave == "img":
            partes.append(t)
            continue
        b, i, u, s, sup, sub, tam = clave
        h = escape(t, quote=False)
        for activo, tag in ((b, "b"), (i, "i"), (u, "u"), (s, "s"), (sup, "sup"), (sub, "sub")):
            if activo:
                h = f"<{tag}>{h}</{tag}>"
        if tam:
            h = f'<span style="font-size:{tam}pt">{h}</span>'
        partes.append(h)
    # \n (salto de línea manual de Word) → <br>
    contenido = "".join(partes).replace("\n", "<br>")

    estilo = []
    al = _alineacion_par(par)
    if al:
        estilo.append(f"text-align:{al}")
    lista = _tipo_lista(par, doc)
    izq = _indent_par(par, "left_indent")
    if izq and izq > 0 and not lista:
        estilo.append(f"margin-left:{izq}pt")
    pri = _indent_par(par, "first_line_indent")
    if pri and not lista:
        estilo.append(f"text-indent:{pri}pt")
    antes, despues, interl = _espaciado(par, doc)
    if antes:
        estilo.append(f"margin-top:{round(antes, 1)}pt")
    if despues:
        estilo.append(f"margin-bottom:{round(despues, 1)}pt")
    if isinstance(interl, tuple):
        estilo.append(f"line-height:{interl[1]}pt")
    elif interl and abs(interl - 1.0) > 0.01:
        estilo.append(f"line-height:{round(1.15 * interl, 2)}")
    st = f' style="{";".join(estilo)}"' if estilo else ""
    dp = f' data-p="{idx}"' if idx is not None else ""

    nombre_estilo = (par.style.name if par.style is not None else "") or ""
    tag = "p"
    if re.match(r"^(Heading|Título|Titulo)\s*1$", nombre_estilo) or nombre_estilo in ("Title", "Título"):
        tag = "h2"
    elif re.match(r"^(Heading|Título|Titulo)\s*[2-6]$", nombre_estilo):
        tag = "h3"
    if lista:
        tag = "li"
    return lista, f"<{tag}{dp}{st}>{contenido or '<br>'}</{tag}>"


def _tabla_html(tbl_el, doc, part, previews, data_t=None) -> str:
    from docx.table import Table
    tabla = Table(tbl_el, _Padre(part))
    dt = f' data-t="{data_t}"' if data_t is not None else ""
    out = [f"<table{dt}><tbody>"]
    for fila in tabla.rows:
        out.append("<tr>")
        vistos = set()
        for celda in fila.cells:
            if id(celda._tc) in vistos:  # celdas combinadas se repiten
                continue
            vistos.add(id(celda._tc))
            pars = "".join(_par_a_html(p._p, doc, part, None, previews)[1] for p in celda.paragraphs)
            out.append(f"<td>{pars}</td>")
        out.append("</tr>")
    out.append("</tbody></table>")
    return "".join(out)


def _tiene_dibujos(el) -> bool:
    return any(True for _ in el.iter(_qn("w:drawing"))) or any(True for _ in el.iter(_qn("w:pict")))


def _bloque_html(el, doc, part, idx, previews) -> tuple[str | None, str]:
    """(tipo_lista, html) de un elemento del cuerpo del Word."""
    if el.tag == _qn("w:p"):
        return _par_a_html(el, doc, part, idx, previews)
    if el.tag == _qn("w:tbl"):
        if idx is not None and _tiene_dibujos(el):
            # Tabla con logos (ej. membrete armado con tabla): va intacta
            return None, (f'<div data-b="{idx}" contenteditable="false" class="bloque-fijo">'
                          f'{_tabla_html(el, doc, part, previews)}</div>')
        return None, _tabla_html(el, doc, part, previews, idx)
    if el.tag == _qn("w:sdt"):
        cont = el.find(_qn("w:sdtContent"))
        interior = "".join(_bloque_html(h, doc, part, None, previews)[1] for h in cont.iterchildren()) \
            if cont is not None else ""
        if idx is None:
            return None, interior
        return None, f'<div data-b="{idx}" contenteditable="false" class="bloque-fijo">{interior or "<p><br></p>"}</div>'
    return None, ""


def _elementos_html(elementos, doc, part, con_indice, previews) -> str:
    out = []
    lista_abierta = None
    for idx, el in elementos:
        if el.tag == _qn("w:sectPr"):
            continue
        lista, h = _bloque_html(el, doc, part, idx if con_indice else None, previews)
        if not h:
            continue
        if lista != lista_abierta:
            if lista_abierta:
                out.append(f"</{lista_abierta}>")
            if lista:
                out.append(f"<{lista}>")
            lista_abierta = lista
        out.append(h)
    if lista_abierta:
        out.append(f"</{lista_abierta}>")
    return "".join(out)


def docx_a_html(datos: bytes, previews: bool = True) -> str:
    """Cuerpo de un .docx como HTML editable (en orden, con referencia al original)."""
    from docx import Document
    doc = Document(io.BytesIO(datos))
    return _elementos_html(enumerate(doc.element.body.iterchildren()), doc, doc.part, True, previews)


def _membrete_parte(doc, parte):
    """Encabezado o pie en uso (el de la última sección que lo define)."""
    try:
        for s in reversed(list(doc.sections)):
            cand = s.header if parte == "header" else s.footer
            if not cand.is_linked_to_previous:
                return cand
    except Exception:
        pass
    return None


def membrete_html(datos: bytes | None) -> dict:
    """Encabezado y pie del Word para mostrarlos (no se editan)."""
    if not datos:
        return {"encabezado": "", "pie": ""}
    from docx import Document
    try:
        doc = Document(io.BytesIO(datos))
    except Exception:
        return {"encabezado": "", "pie": ""}
    out = {}
    for parte, clave in (("header", "encabezado"), ("footer", "pie")):
        hf = _membrete_parte(doc, parte)
        html = ""
        if hf is not None:
            try:
                html = _elementos_html(((None, el) for el in hf._element.iterchildren()), doc, hf.part, False, True)
            except Exception as e:
                print(f"[!] No se pudo leer el {parte}: {e}")
        # Si no tiene texto ni imágenes, no se muestra
        if "<img" not in html and not re.sub(r"<[^>]+>", "", html).strip():
            html = ""
        out[clave] = html
    return out


def docx_valido(datos: bytes) -> bool:
    try:
        from docx import Document
        Document(io.BytesIO(datos))
        return True
    except Exception:
        return False


# ═══════════════════════════════════════════════════════════════
# Variables @ de las plantillas
# ═══════════════════════════════════════════════════════════════

_VAR_RE = re.compile(r"@(n[uú]mero|car[aá]tula|juzgado|fecha|mes)(?![\wáéíóúñÁÉÍÓÚÑ])", re.IGNORECASE)


def _clave_var(txt: str) -> str:
    return txt.lower().replace("ú", "u").replace("á", "a")


def _reemplazar_en_parrafo(p_el, valores: dict) -> int:
    ts = []
    for r in _iter_runs(p_el):
        ts.extend(r.findall(_qn("w:t")))
    if not ts:
        return 0
    completo = "".join(t.text or "" for t in ts)
    coincidencias = list(_VAR_RE.finditer(completo))
    if not coincidencias:
        return 0
    # Posición de cada w:t dentro del texto completo
    tramos, pos = [], 0
    for t in ts:
        n = len(t.text or "")
        tramos.append([t, pos, pos + n])
        pos += n
    for m in reversed(coincidencias):
        ini, fin = m.span()
        valor = valores.get(_clave_var(m.group(1)), m.group(0))
        primero = True
        for t, a, b in tramos:
            if b <= ini or a >= fin:
                continue
            txt = t.text or ""
            if primero:
                t.text = txt[:max(ini - a, 0)] + valor + (txt[fin - a:] if fin < b else "")
                primero = False
            else:
                t.text = txt[fin - a:] if fin < b else ""
            t.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
        # recalcular posiciones para los reemplazos anteriores (van de atrás hacia adelante)
        pos = 0
        for tr in tramos:
            n = len(tr[0].text or "")
            tr[1], tr[2] = pos, pos + n
            pos += n
    return len(coincidencias)


def reemplazar_variables(datos: bytes, valores: dict) -> bytes:
    """Completa @numero, @caratula, @juzgado, @fecha y @mes en todo el Word
    (cuerpo, tablas, encabezado y pie), respetando el formato de cada palabra."""
    from docx import Document
    doc = Document(io.BytesIO(datos))
    total = 0
    raices = [doc.element.body]
    for rel in doc.part.rels.values():
        if rel.reltype.endswith("/header") or rel.reltype.endswith("/footer"):
            try:
                raices.append(rel.target_part.element)
            except Exception:
                pass
    for raiz in raices:
        for p in raiz.iter(_qn("w:p")):
            total += _reemplazar_en_parrafo(p, valores)
    if not total:
        return datos
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


# ═══════════════════════════════════════════════════════════════
# Formato dominante del Word base (para lo que se escribe nuevo)
# ═══════════════════════════════════════════════════════════════

def _formato_base(doc) -> dict:
    fuentes, tamanos, interl, despues, antes = Counter(), Counter(), Counter(), Counter(), Counter()
    for par in doc.paragraphs:
        largo = len(par.text.strip())
        if not largo:
            continue
        pf = par.paragraph_format
        if pf.line_spacing is not None:
            interl[pf.line_spacing if isinstance(pf.line_spacing, float) else ("pt", pf.line_spacing.pt)] += largo
        if pf.space_after is not None:
            despues[pf.space_after.pt] += largo
        if pf.space_before is not None:
            antes[pf.space_before.pt] += largo
        for r in par.runs:
            n = len(r.text)
            if r.font.name:
                fuentes[r.font.name] += n
            if r.font.size:
                tamanos[r.font.size.pt] += n
    normal = None
    try:
        normal = doc.styles["Normal"]
    except KeyError:
        pass
    f = {
        "fuente": fuentes.most_common(1)[0][0] if fuentes else None,
        "tamano": tamanos.most_common(1)[0][0] if tamanos else None,
        "interlineado": interl.most_common(1)[0][0] if interl else None,
        "despues": despues.most_common(1)[0][0] if despues else None,
        "antes": antes.most_common(1)[0][0] if antes else None,
    }
    f["fuente_efectiva"] = f["fuente"] or (normal.font.name if normal is not None and normal.font.name else None) \
        or _fuente_por_defecto(doc)
    f["tamano_efectivo"] = f["tamano"] or _defaults(doc)["tam"] or 12
    if f["interlineado"] is None and normal is not None:
        ls = normal.paragraph_format.line_spacing
        if isinstance(ls, float):
            f["interlineado_efectivo"] = ls
    f.setdefault("interlineado_efectivo", f["interlineado"] if isinstance(f["interlineado"], float) else 1.0)
    return f


def _fuente_por_defecto(doc):
    try:
        v = doc.styles.element.xpath("./w:docDefaults/w:rPrDefault/w:rPr/w:rFonts/@w:ascii")
        return v[0] if v else None
    except Exception:
        return None


# ═══════════════════════════════════════════════════════════════
# HTML → Word
# ═══════════════════════════════════════════════════════════════

def _doc_en_blanco():
    """Hoja en blanco: A4, márgenes judiciales, Times New Roman 12, interlineado 1,5."""
    from docx import Document
    from docx.shared import Cm, Pt
    doc = Document()
    s = doc.sections[0]
    s.page_width, s.page_height = Cm(21), Cm(29.7)
    s.left_margin, s.right_margin = Cm(3), Cm(2)
    s.top_margin, s.bottom_margin = Cm(2.5), Cm(2.5)
    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(12)
    normal.paragraph_format.line_spacing = 1.5
    normal.paragraph_format.space_after = Pt(0)
    for hijo in list(doc.element.body.iterchildren()):
        if hijo.tag != _qn("w:sectPr"):
            doc.element.body.remove(hijo)
    return doc, {"fuente": None, "tamano": None, "interlineado": None, "despues": None, "antes": None,
                 "fuente_efectiva": "Times New Roman", "tamano_efectivo": 12, "interlineado_efectivo": 1.5}


def _aplicar_par(par, b, fmt):
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Pt
    pf = par.paragraph_format
    al = b.get("align")
    if al:
        pf.alignment = {"center": WD_ALIGN_PARAGRAPH.CENTER, "right": WD_ALIGN_PARAGRAPH.RIGHT,
                        "justify": WD_ALIGN_PARAGRAPH.JUSTIFY, "left": WD_ALIGN_PARAGRAPH.LEFT}[al]
    if b.get("izq"):
        pf.left_indent = Pt(b["izq"])
    if b.get("primera") is not None:
        pf.first_line_indent = Pt(b["primera"])
    il = fmt.get("interlineado")
    if il is not None:
        pf.line_spacing = il if isinstance(il, float) else Pt(il[1])
    if fmt.get("despues") is not None:
        pf.space_after = Pt(fmt["despues"])
    if fmt.get("antes") is not None:
        pf.space_before = Pt(fmt["antes"])


def _formatear_run(run, r, b):
    from docx.shared import Pt
    run.bold = bool(r.get("b") or b.get("titulo"))
    run.italic = bool(r.get("i"))
    run.underline = bool(r.get("u"))
    run.font.strike = bool(r.get("s"))
    run.font.superscript = True if r.get("sup") else None
    run.font.subscript = True if r.get("sub") else None
    if r.get("sz"):
        run.font.size = Pt(r["sz"])


def _aplicar_runs(par, b, fmt, imagenes=None):
    from docx.shared import Pt
    for r in b["runs"]:
        if "img" in r:
            orig = (imagenes or {}).get(r["img"])
            if orig is not None:
                par._p.append(copy.deepcopy(orig))
            continue
        if r.get("br"):
            par.add_run().add_break()
            continue
        run = par.add_run(r["t"])
        _formatear_run(run, r, b)
        if fmt.get("fuente"):
            run.font.name = fmt["fuente"]
            rpr = run._r.get_or_add_rPr()
            rfonts = rpr.find(_qn("w:rFonts"))
            if rfonts is not None:
                rfonts.set(_qn("w:eastAsia"), fmt["fuente"])
                rfonts.set(_qn("w:cs"), fmt["fuente"])
        if fmt.get("tamano") and not r.get("sz"):
            run.font.size = Pt(fmt["tamano"])


def _estilo_si_existe(doc, *nombres):
    for n in nombres:
        try:
            return doc.styles[n]
        except KeyError:
            continue
    return None


def _agregar_bloques(contenedor, bloques, fmt, doc, primer_par=None, imagenes=None):
    """Arma bloques desde cero (hoja en blanco, celdas de tablas editadas)."""
    for b in bloques:
        if b["tipo"] == "tabla":
            _agregar_tabla(contenedor, b, fmt, doc, imagenes)
            continue
        if b["tipo"] == "fijo":
            _agregar_bloques(contenedor, b["bloques"], fmt, doc, imagenes=imagenes)
            continue
        if primer_par is not None:
            par, primer_par = primer_par, None
        else:
            par = contenedor.add_paragraph()
        if b.get("lista"):
            estilo = _estilo_si_existe(doc, "List Bullet", "Lista con viñetas")
            if estilo is not None and b["lista"] == "ul":
                par.style = estilo
            else:
                pref = "•\t" if b["lista"] == "ul" else f"{b.get('numero', 1)}.\t"
                b = dict(b, runs=[{"t": pref}] + b["runs"])
                b["izq"] = (b.get("izq") or 0) + 18
                b["primera"] = -18
        if b.get("titulo") == 1:
            est = _estilo_si_existe(doc, "Heading 1", "Título 1")
            if est is not None:
                par.style = est
        elif b.get("titulo") == 2:
            est = _estilo_si_existe(doc, "Heading 2", "Título 2")
            if est is not None:
                par.style = est
        _aplicar_par(par, b, fmt)
        _aplicar_runs(par, b, fmt, imagenes)


def _agregar_tabla(contenedor, b, fmt, doc, imagenes=None):
    ncols = max(sum(c["colspan"] for c in fila) for fila in b["filas"])
    tabla = contenedor.add_table(rows=len(b["filas"]), cols=ncols)
    est = _estilo_si_existe(doc, "Table Grid", "Tabla con cuadrícula")
    if est is not None:
        try:
            tabla.style = est
        except Exception:
            pass
    else:
        _bordes_tabla(tabla)
    for i, fila in enumerate(b["filas"]):
        col = 0
        for c in fila:
            celda = tabla.cell(i, col)
            if c["colspan"] > 1 and col + c["colspan"] - 1 < ncols:
                celda = celda.merge(tabla.cell(i, col + c["colspan"] - 1))
            _agregar_bloques(celda, c["bloques"] or [{"tipo": "p", "runs": []}], fmt, doc,
                             primer_par=celda.paragraphs[0], imagenes=imagenes)
            col += c["colspan"]
    return tabla


def _bordes_tabla(tabla):
    from docx.oxml import OxmlElement
    tbl_pr = tabla._tbl.tblPr
    bordes = OxmlElement("w:tblBorders")
    for lado in ("top", "left", "bottom", "right", "insideH", "insideV"):
        e = OxmlElement(f"w:{lado}")
        e.set(_qn("w:val"), "single")
        e.set(_qn("w:sz"), "4")
        e.set(_qn("w:color"), "000000")
        bordes.append(e)
    tbl_pr.append(bordes)


def _primer_rpr(p_el):
    for r in _iter_runs(p_el):
        if _es_especial(r):
            continue
        if r.find(_qn("w:t")) is not None:
            rpr = r.find(_qn("w:rPr"))
            return copy.deepcopy(rpr) if rpr is not None else None
    # párrafo vacío: la marca de párrafo guarda el formato de letra
    ppr = p_el.find(_qn("w:pPr"))
    if ppr is not None:
        mark = ppr.find(_qn("w:rPr"))
        if mark is not None:
            from docx.oxml import OxmlElement
            rpr = OxmlElement("w:rPr")
            for h in mark.iterchildren():
                if h.tag not in (_qn("w:ins"), _qn("w:del"), _qn("w:moveFrom"), _qn("w:moveTo")):
                    rpr.append(copy.deepcopy(h))
            return rpr
    return None


def _nuevo_parrafo(b, ppr, rpr, imagenes, doc, conservar_seccion: bool):
    """Párrafo corregido: mismo formato de párrafo y de letra que el original."""
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.shared import Pt
    from docx.text.paragraph import Paragraph
    from docx.text.run import Run
    p = OxmlElement("w:p")
    tiene_num = False
    if ppr is not None:
        ppr2 = copy.deepcopy(ppr)
        for quitar in ([] if conservar_seccion else [_qn("w:sectPr")]) + [_qn("w:rPrChange"), _qn("w:pPrChange")]:
            for x in ppr2.findall(quitar):
                ppr2.remove(x)
        num = ppr2.find(_qn("w:numPr"))
        if num is not None and not b.get("lista"):
            ppr2.remove(num)
        tiene_num = ppr2.find(_qn("w:numPr")) is not None
        p.append(ppr2)
    par = Paragraph(p, None)
    pf = par.paragraph_format
    pf.alignment = {"center": WD_ALIGN_PARAGRAPH.CENTER, "right": WD_ALIGN_PARAGRAPH.RIGHT,
                    "justify": WD_ALIGN_PARAGRAPH.JUSTIFY}.get(b.get("align"), WD_ALIGN_PARAGRAPH.LEFT)
    if not tiene_num:
        pf.left_indent = Pt(b.get("izq") or 0)
        pf.first_line_indent = Pt(b.get("primera") or 0)
    runs = b["runs"]
    if b.get("lista") and not tiene_num:
        pref = "•\t" if b["lista"] == "ul" else f"{b.get('numero', 1)}.\t"
        runs = [{"t": pref}] + runs
        pf.left_indent = Pt((b.get("izq") or 0) + 18)
        pf.first_line_indent = Pt(-18)
    for r in runs:
        if "img" in r:
            orig = imagenes.get(r["img"])
            if orig is not None:
                p.append(copy.deepcopy(orig))
            continue
        rr = OxmlElement("w:r")
        if rpr is not None:
            rr.append(copy.deepcopy(rpr))
        p.append(rr)
        run = Run(rr, par)
        if r.get("br"):
            run.add_break()
            continue
        run.text = r["t"]
        _formatear_run(run, r, b)
    return p


def armar_docx(base: bytes | None, html: str) -> bytes:
    from docx import Document
    bloques = html_a_bloques(html) or [{"tipo": "p", "runs": []}]
    if not base:
        doc, fmt = _doc_en_blanco()
        _agregar_bloques(doc, bloques, fmt, doc)
        if bloques[-1]["tipo"] == "tabla":
            doc.add_paragraph()
        buf = io.BytesIO()
        doc.save(buf)
        return buf.getvalue()

    doc = Document(io.BytesIO(base))
    fmt = _formato_base(doc)
    cuerpo = doc.element.body
    originales = list(cuerpo.iterchildren())
    seccion = cuerpo.find(_qn("w:sectPr"))

    # Imágenes originales por clave "párrafo:run" (se copian donde sigan estando)
    imagenes = {}
    for i, el in enumerate(originales):
        if el.tag == _qn("w:p"):
            for j, r in enumerate(_iter_runs(el)):
                if _es_especial(r):
                    imagenes[f"{i}:{j}"] = r

    firmas = {}

    def firma_original(i):
        if i not in firmas:
            el = originales[i]
            lista, h = _bloque_html(el, doc, doc.part, i, False)
            if lista:
                h = f"<{lista}>{h}</{lista}>"
            bs = html_a_bloques(h)
            firmas[i] = _firma(bs[0]) if bs else None
        return firmas[i]

    for hijo in originales:
        if hijo is not seccion:
            cuerpo.remove(hijo)

    def agregar(el):
        if seccion is not None:
            seccion.addprevious(el)
        else:
            cuerpo.append(el)

    def valido(i, tag):
        return i is not None and 0 <= i < len(originales) and originales[i].tag == _qn(tag)

    usados = set()
    formato_previo = None  # (pPr, rPr) del último párrafo original: para lo que se escribe nuevo
    for b in bloques:
        o = b.get("origen")
        if b["tipo"] == "fijo":
            if o is not None and 0 <= o < len(originales) and o not in usados:
                agregar(copy.deepcopy(originales[o]))
                usados.add(o)
            continue
        if b["tipo"] == "tabla":
            if valido(o, "w:tbl") and o not in usados and _firma(b) == firma_original(o):
                agregar(copy.deepcopy(originales[o]))
                usados.add(o)
            else:
                _agregar_tabla(doc, b, fmt, doc, imagenes)
            continue
        if valido(o, "w:p"):
            orig = originales[o]
            primera_vez = o not in usados
            usados.add(o)
            if primera_vez and _firma(b) == firma_original(o):
                agregar(copy.deepcopy(orig))
            else:
                agregar(_nuevo_parrafo(b, orig.find(_qn("w:pPr")), _primer_rpr(orig), imagenes, doc,
                                       conservar_seccion=primera_vez))
            formato_previo = (orig.find(_qn("w:pPr")), _primer_rpr(orig))
            continue
        if formato_previo is not None:
            agregar(_nuevo_parrafo(b, formato_previo[0], formato_previo[1], imagenes, doc, conservar_seccion=False))
        else:
            _agregar_bloques(doc, [b], fmt, doc, imagenes=imagenes)

    # Word exige que el documento no termine en tabla
    ultimo = [h for h in cuerpo.iterchildren() if h is not seccion]
    if ultimo and ultimo[-1].tag == _qn("w:tbl"):
        doc.add_paragraph()
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


# ═══════════════════════════════════════════════════════════════
# HTML → PDF
# ═══════════════════════════════════════════════════════════════

def armar_pdf(base: bytes | None, html: str) -> bytes:
    """PDF fiel: el Word armado convertido por LibreOffice. Si no hay LibreOffice
    (ej. en una PC sin instalar), se arma uno aproximado con reportlab."""
    docx = armar_docx(base, html)
    pdf = convertir_libreoffice(docx, ".docx", "pdf")
    if pdf:
        return pdf
    return _pdf_reportlab(base, html)


_SANS = ("arial", "helvetica", "calibri", "verdana", "tahoma", "segoe", "trebuchet", "century gothic", "aptos")


def _familia(fuente: str | None) -> tuple[str, str, str, str]:
    f = (fuente or "").lower()
    if any(s in f for s in _SANS):
        return "Helvetica", "Helvetica-Bold", "Helvetica-Oblique", "Helvetica-BoldOblique"
    if "courier" in f:
        return "Courier", "Courier-Bold", "Courier-Oblique", "Courier-BoldOblique"
    return "Times-Roman", "Times-Bold", "Times-Italic", "Times-BoldItalic"


def _registrar_familia(fams):
    from reportlab.lib.fonts import addMapping
    n, b, i, bi = fams
    addMapping(n, 0, 0, n)
    addMapping(n, 1, 0, b)
    addMapping(n, 0, 1, i)
    addMapping(n, 1, 1, bi)


def _texto_rl(t: str) -> str:
    t = t.replace("\t", " " * 8).replace(" ", "  ").replace(" ", " ")
    t = re.sub(r"  +", lambda m: " " * len(m.group(0)), t)
    return escape(t, quote=False)


def _runs_rl(b) -> str:
    out = []
    for r in b["runs"]:
        if "img" in r:
            continue
        if r.get("br"):
            out.append("<br/>")
            continue
        h = _texto_rl(r["t"])
        for k, tag in (("b", "b"), ("i", "i"), ("u", "u"), ("s", "strike"), ("sup", "super"), ("sub", "sub")):
            if r.get(k):
                h = f"<{tag}>{h}</{tag}>"
        if r.get("sz"):
            h = f'<font size="{r["sz"]}">{h}</font>'
        out.append(h)
    txt = "".join(out)
    if b.get("titulo"):
        txt = f"<b>{txt}</b>"
    return txt


def _flowables(bloques, fmt, ancho, fams):
    from reportlab.lib.enums import TA_CENTER, TA_JUSTIFY, TA_LEFT, TA_RIGHT
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.platypus import Paragraph, Table, TableStyle
    from reportlab.lib import colors

    tam = float(fmt.get("tamano_efectivo") or 12)
    alinear = {"center": TA_CENTER, "right": TA_RIGHT, "justify": TA_JUSTIFY, "left": TA_LEFT}

    out = []
    for b in bloques:
        if b["tipo"] == "fijo":
            out.extend(_flowables(b["bloques"], fmt, ancho, fams))
            continue
        if b["tipo"] == "tabla":
            ncols = max(sum(c["colspan"] for c in fila) for fila in b["filas"])
            datos, spans = [], []
            for i, fila in enumerate(b["filas"]):
                fila_rl, col = [], 0
                for c in fila:
                    fila_rl.append(_flowables(c["bloques"], fmt, ancho / ncols, fams) or "")
                    if c["colspan"] > 1:
                        spans.append(("SPAN", (col, i), (col + c["colspan"] - 1, i)))
                        fila_rl.extend([""] * (c["colspan"] - 1))
                    col += c["colspan"]
                fila_rl.extend([""] * (ncols - len(fila_rl)))
                datos.append(fila_rl)
            t = Table(datos, colWidths=[ancho / ncols] * ncols)
            t.setStyle(TableStyle([("GRID", (0, 0), (-1, -1), 0.5, colors.black),
                                   ("VALIGN", (0, 0), (-1, -1), "TOP")] + spans))
            out.append(t)
            continue
        il = b.get("interl", fmt.get("interlineado"))
        if isinstance(il, tuple):
            leading = il[1]
        else:
            leading = tam * 1.17 * float(il if isinstance(il, float) else fmt.get("interlineado_efectivo") or 1.0)
        izq = b.get("izq") or 0
        pri = b.get("primera") or 0
        bullet = None
        if b.get("lista"):
            izq += 18
            bullet = "•" if b["lista"] == "ul" else f"{b.get('numero', 1)}."
        tam_b = tam * (1.25 if b.get("titulo") == 1 else 1.1 if b.get("titulo") == 2 else 1)
        estilo = ParagraphStyle(
            "p", fontName=fams[0], fontSize=tam_b, leading=leading * tam_b / tam,
            alignment=alinear.get(b.get("align") or "left", TA_LEFT),
            leftIndent=izq, firstLineIndent=pri,
            spaceAfter=b.get("despues", fmt.get("despues") or 0), spaceBefore=b.get("antes", fmt.get("antes") or 0),
            bulletIndent=izq - 14 if bullet else 0, bulletFontName=fams[0], bulletFontSize=tam_b,
        )
        txt = _runs_rl(b) or " "
        try:
            out.append(Paragraph(txt, estilo, bulletText=bullet))
        except Exception:
            out.append(Paragraph(escape("".join(r.get("t", " ") for r in b["runs"])) or " ", estilo))
    return out


def _encabezado_pie(doc, parte, fmt, fams):
    """Flowables del encabezado o pie (texto + imágenes)."""
    from docx.table import Table as DTable
    from docx.text.paragraph import Paragraph as DPar
    from reportlab.platypus import Image as RLImage
    hf = _membrete_parte(doc, parte)
    if hf is None:
        return []
    sec = doc.sections[-1]
    out = []
    ancho_util = (sec.page_width - sec.left_margin - sec.right_margin) / 12700
    plano = dict(fmt, despues=0, antes=0, interlineado=1.0)

    def de_parrafo(el):
        par = DPar(el, _Padre(hf.part))
        for blip in el.iter(f"{{{_NS_A}}}blip"):
            rid = blip.get(f"{{{_NS_R}}}embed")
            try:
                blob = hf.part.related_parts[rid].blob
                ext = next(blip.iterancestors(f"{{{_NS_WP}}}inline", f"{{{_NS_WP}}}anchor")).find(f"{{{_NS_WP}}}extent")
                w = int(ext.get("cx")) / 12700 if ext is not None else 100
                h = int(ext.get("cy")) / 12700 if ext is not None else 50
                if w > ancho_util:
                    h, w = h * ancho_util / w, ancho_util
                img = RLImage(io.BytesIO(blob), width=w, height=h)
                img.hAlign = {"center": "CENTER", "right": "RIGHT"}.get(_alineacion_par(par), "LEFT")
                out.append(img)
            except Exception:
                continue
        if par.text.strip():
            _, h = _par_a_html(el, doc, hf.part, None, False)
            out.extend(_flowables(html_a_bloques(h), plano, ancho_util, fams))

    for el in hf._element.iterchildren():
        if el.tag == _qn("w:p"):
            de_parrafo(el)
        elif el.tag == _qn("w:tbl"):
            for fila in DTable(el, _Padre(hf.part)).rows:
                for celda in fila.cells:
                    for p in celda.paragraphs:
                        de_parrafo(p._p)
    return out


def _pdf_reportlab(base: bytes | None, html: str) -> bytes:
    from reportlab.platypus import SimpleDocTemplate, Frame, KeepInFrame
    from reportlab.lib.pagesizes import A4

    pagina, margenes = A4, (85, 57, 71, 71)  # izq, der, arriba, abajo (pt)
    enc, pie = [], []
    fmt = {"fuente_efectiva": "Times New Roman", "tamano_efectivo": 12, "interlineado_efectivo": 1.5,
           "interlineado": 1.5}
    dist_enc = dist_pie = 35
    doc = None
    if base:
        from docx import Document
        doc = Document(io.BytesIO(base))
        fmt = _formato_base(doc)
        s = doc.sections[-1]
        try:
            pagina = (s.page_width.pt, s.page_height.pt)
            margenes = (s.left_margin.pt, s.right_margin.pt, s.top_margin.pt, s.bottom_margin.pt)
            dist_enc = s.header_distance.pt if s.header_distance is not None else 35
            dist_pie = s.footer_distance.pt if s.footer_distance is not None else 35
        except Exception:
            pass
    fams = _familia(fmt.get("fuente_efectiva"))
    _registrar_familia(fams)
    if doc is not None:
        enc = _encabezado_pie(doc, "header", fmt, fams)
        pie = _encabezado_pie(doc, "footer", fmt, fams)

    izq, der, arriba, abajo = margenes
    ancho = pagina[0] - izq - der

    # Como en Word: si el membrete no entra en el margen, empuja el cuerpo
    def _alto(fl):
        total = 0
        for x in fl:
            try:
                total += x.wrap(ancho, 10000)[1] + x.getSpaceBefore() + x.getSpaceAfter()
            except Exception:
                pass
        return total
    if enc:
        arriba = max(arriba, dist_enc + _alto(enc) + 8)
    if pie:
        abajo = max(abajo, dist_pie + _alto(pie) + 8)

    def en_pagina(canvas, _doc):
        canvas.saveState()
        if enc:
            alto = max(arriba - dist_enc - 2, 20)
            f = Frame(izq, pagina[1] - dist_enc - alto, ancho, alto, 0, 0, 0, 0, showBoundary=0)
            f.addFromList([KeepInFrame(ancho, alto, list(enc), mode="shrink")], canvas)
        if pie:
            alto = max(abajo - dist_pie - 2, 20)
            f = Frame(izq, dist_pie, ancho, alto, 0, 0, 0, 0, showBoundary=0)
            f.addFromList([KeepInFrame(ancho, alto, list(pie), mode="shrink")], canvas)
        canvas.restoreState()

    buf = io.BytesIO()
    pdf = SimpleDocTemplate(buf, pagesize=pagina, leftMargin=izq, rightMargin=der,
                            topMargin=arriba, bottomMargin=abajo, title="Proyecto de dictamen")
    cuerpo = _flowables(html_a_bloques(html) or [{"tipo": "p", "runs": []}], fmt, ancho, fams)
    pdf.build(cuerpo, onFirstPage=en_pagina, onLaterPages=en_pagina)
    return buf.getvalue()
