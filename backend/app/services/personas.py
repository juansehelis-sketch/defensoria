"""
Lectura de personas (nombre, DNI, fecha de nacimiento) desde un PDF o Word.

Sirve para cargar los defendidos de un expediente sin tipear: se sube la vista,
la demanda o un escrito, y se proponen las personas que aparecen con su DNI o
su fecha de nacimiento. Nada se guarda solo: la persona revisa y elige.

Se basa en cómo se escriben estos datos en los escritos judiciales:
  "PÉREZ, Juan Carlos, DNI N° 45.123.456"
  "Juan Pérez (D.N.I. 45123456)"
  "la representación de Vicente Vara, nacido el 8 de mayo de 2012"
  "fecha de nacimiento: 08/05/2012"
"""

import io
import re
import unicodedata
from datetime import date

_MESES = {"enero": 1, "febrero": 2, "marzo": 3, "abril": 4, "mayo": 5, "junio": 6, "julio": 7,
          "agosto": 8, "septiembre": 9, "setiembre": 9, "octubre": 10, "noviembre": 11, "diciembre": 12}

_LETRA = "A-Za-zÁÉÍÓÚÜÑáéíóúüñ"
_PAL_MAY = rf"[A-ZÁÉÍÓÚÜÑ][{_LETRA}'\-]+"
_CONECTOR = r"(?:de|del|la|las|los|y|da|di|van|von)"

# DNI: "DNI N° 45.123.456", "D.N.I. 45123456", "documento nacional de identidad n° 45.123.456"
_DNI_RE = re.compile(
    r"(?:D\.?\s?N\.?\s?I\.?|documento(?:\s+nacional)?(?:\s+de\s+identidad)?)\s*"
    r"(?:n(?:ro|úmero|umero)?\.?\s*[°ºo]?\s*|[°º]\s*)?[:.]?\s*"
    r"(\d{1,2}[.\s]?\d{3}[.\s]?\d{3})\b",
    re.IGNORECASE,
)

_FECHA_NUM = r"(\d{1,2})[/\-.](\d{1,2})[/\-.](\d{2,4})"
_FECHA_LETRAS = r"(\d{1,2})(?:º|°)?\s+de\s+([a-zA-Z]+)\s+(?:de|del)\s+(\d{4})"
_NAC_RE = re.compile(
    rf"(?:nacid[oa]s?\s+el(?:\s+d[ií]a)?|fecha\s+de\s+nacimiento|f\.\s*nac\.?|nac\.)\s*[:\-]?\s*"
    rf"(?:{_FECHA_NUM}|{_FECHA_LETRAS})",
    re.IGNORECASE,
)

# Palabras que indican que NO es un defendido (profesionales, funcionarios)
_NO_PERSONA = re.compile(
    r"\b(dr|dra|doctor|doctora|juez|jueza|fiscal|defensor[a]?|secretari[oa]|lic|licenciad[oa]|asesor[a]?|perit[oa]|letrad[oa]|abogad[oa])\b\.?\s*$",
    re.IGNORECASE,
)
_PALABRAS_NO_NOMBRE = {
    "señor", "señora", "juez", "jueza", "sr", "sra", "vs", "autos", "expte", "expediente", "civil", "juzgado",
    "defensoría", "defensoria", "ministerio", "público", "publico", "tribunal", "cámara", "camara", "nación",
    "nacion", "buenos", "aires", "ciudad", "autónoma", "autonoma", "república", "argentina", "ley", "art", "arts",
    "código", "codigo", "ccyc", "que", "el", "la", "los", "las", "su", "sus", "en", "con", "por", "para", "a",
    "asumo", "representación", "representacion", "menor", "niño", "niña", "hijo", "hija", "nna", "titular",
    "identidad", "documento", "nacional", "nacido", "nacida", "quien", "cuyo", "cuya", "sobre", "se", "vista",
    "dni", "d.n.i", "nro", "n", "fs", "fojas", "cuil", "cuit", "sres", "sras", "señores", "doña", "don",
}


def extraer_texto(datos: bytes, nombre: str) -> str:
    """Texto de un PDF o de un Word."""
    n = (nombre or "").lower()
    if n.endswith(".pdf"):
        import pdfplumber
        partes = []
        with pdfplumber.open(io.BytesIO(datos)) as pdf:
            for pag in pdf.pages[:60]:
                partes.append(pag.extract_text() or "")
        return "\n".join(partes)
    if n.endswith(".docx"):
        from docx import Document
        doc = Document(io.BytesIO(datos))
        partes = [p.text for p in doc.paragraphs]
        for t in doc.tables:
            for fila in t.rows:
                for c in fila.cells:
                    partes.append(c.text)
        return "\n".join(partes)
    raise ValueError("El archivo tiene que ser PDF o Word (.docx)")


def _sin_tildes(t: str) -> str:
    return "".join(c for c in unicodedata.normalize("NFD", t) if unicodedata.category(c) != "Mn")


def normalizar_dni(dni: str | None) -> str:
    return re.sub(r"\D", "", dni or "")


def formatear_dni(dni: str) -> str:
    d = normalizar_dni(dni)
    if len(d) < 7:
        return d
    return f"{int(d[:-6])}.{d[-6:-3]}.{d[-3:]}"


def _fecha(m, desde: int) -> date | None:
    g = m.groups()[desde:desde + 6]
    try:
        if g[0]:
            d, mes, a = int(g[0]), int(g[1]), int(g[2])
            if a < 100:
                a += 2000 if a < 30 else 1900
        else:
            d, mes, a = int(g[3]), _MESES.get(_sin_tildes(g[4]).lower()), int(g[5])
            if not mes:
                return None
        f = date(a, mes, d)
        if 1900 <= f.year <= date.today().year:
            return f
    except (TypeError, ValueError):
        return None
    return None


_PROFESIONAL = object()


def _nombre_antes(texto: str, fin: int):
    """Nombre propio que termina justo antes de `fin` (ej. antes de ", DNI").
    Devuelve el nombre, None si no hay, o _PROFESIONAL si es un/a profesional."""
    ventana = texto[max(0, fin - 160):fin]
    tokens = re.findall(rf"[{_LETRA}][{_LETRA}'\-]*\.?|\d[\d.]*|[,()\[\]:;]|\S", ventana)
    # Saltear el conector final: "Juan Pérez, DNI" / "Juan Pérez (DNI" / "titular del DNI"
    while tokens and (tokens[-1] in (",", "(", "-", "–") or
                      _sin_tildes(tokens[-1]).lower() in ("con", "de", "del", "titular")):
        tokens.pop()
    nombre = []
    comas = 0
    for tok in reversed(tokens):
        low = _sin_tildes(tok).lower().rstrip(".")
        if tok == ",":
            if nombre and comas == 0:
                comas += 1
                nombre.insert(0, ",")
                continue
            break
        if re.fullmatch(_CONECTOR, tok):
            if nombre:
                nombre.insert(0, tok)
                continue
            break
        es_nombre = (tok[0].isupper() and re.fullmatch(rf"[{_LETRA}'\-]+", tok) is not None
                     and low not in _PALABRAS_NO_NOMBRE and len(low) > 1)
        if not es_nombre:
            break
        nombre.insert(0, tok)
        if sum(1 for x in nombre if x != "," and not re.fullmatch(_CONECTOR, x)) >= 6:
            break
    while nombre and (nombre[0] == "," or re.fullmatch(_CONECTOR, nombre[0])):
        nombre.pop(0)
    while nombre and (nombre[-1] == "," or re.fullmatch(_CONECTOR, nombre[-1])):
        nombre.pop()
    palabras = [x for x in nombre if x != "," and not re.fullmatch(_CONECTOR, x)]
    if len(palabras) < 2:
        return None
    # ¿Viene precedido de "Dra.", "Lic.", "Juez"...? Entonces es un/a profesional
    txt = " ".join(nombre).replace(" ,", ",")
    idx = ventana.rfind(palabras[0])
    if idx > 0 and _NO_PERSONA.search(ventana[:idx]):
        return _PROFESIONAL
    return re.sub(r"\s+", " ", txt).strip()


def _vinculo(antes: str, fecha_nac) -> str | None:
    """Vínculo según la palabra más cercana ANTES del nombre ("el niño", "su
    progenitora") y, si no hay, según la edad."""
    c = _sin_tildes(antes).lower()
    ultimo = {}
    for tipo, patron in (("NNA", r"(?<![a-z])(nin[oa]s?|menor(es)?|adolescentes?|hij[oa]s?|nna|joven)(?![a-z])"),
                         ("progenitor/a", r"(?<![a-z])(madre|padre|progenitor[a]?|abuel[oa]|tia|tio|guardador[a]?)(?![a-z])")):
        ms = list(re.finditer(patron, c))
        if ms:
            ultimo[tipo] = ms[-1].end()
    if ultimo:
        return max(ultimo, key=ultimo.get)
    if fecha_nac:
        hoy = date.today()
        edad = hoy.year - fecha_nac.year - ((hoy.month, hoy.day) < (fecha_nac.month, fecha_nac.day))
        return "NNA" if edad < 18 else None
    return None


def _clave_nombre(n: str) -> str:
    return re.sub(r"[^a-z]", "", _sin_tildes(n).lower())


def personas_en_texto(texto: str) -> list[dict]:
    """Personas con DNI o fecha de nacimiento que aparecen en el texto."""
    t = re.sub(r"[ \t\xa0]+", " ", texto.replace("\r", ""))
    t = re.sub(r"-\n(?=[a-záéíóúñ])", "", t)   # palabras cortadas al final de renglón
    t = t.replace("\n", " ")
    hallazgos = []

    dnis = list(_DNI_RE.finditer(t))
    for k, m in enumerate(dnis):
        nombre = _nombre_antes(t, m.start())
        if nombre is _PROFESIONAL:
            continue
        # La fecha de nacimiento de ESTA persona: después del DNI y antes del próximo DNI
        hasta = dnis[k + 1].start() if k + 1 < len(dnis) else len(t)
        nac = _NAC_RE.search(t[m.end():min(hasta, m.end() + 160)])
        hallazgos.append({
            "nombre": nombre, "dni": formatear_dni(m.group(1)),
            "fecha_nacimiento": _fecha(nac, 0) if nac else None,
            "pos": m.start(),
        })

    for m in _NAC_RE.finditer(t):
        nombre = _nombre_antes(t, m.start())
        if not nombre or nombre is _PROFESIONAL:
            continue
        hallazgos.append({"nombre": nombre, "dni": None, "fecha_nacimiento": _fecha(m, 0), "pos": m.start()})

    # Unir lo que es la misma persona (mismo DNI, o mismo nombre)
    personas = []
    for h in hallazgos:
        if not h["nombre"] and not h["dni"]:
            continue
        igual = None
        for p in personas:
            if (h["dni"] and p["dni"] == h["dni"]) or \
                    (h["nombre"] and p["nombre"] and _clave_nombre(h["nombre"]) == _clave_nombre(p["nombre"])):
                igual = p
                break
        if igual is None:
            contexto = t[max(0, h["pos"] - 140):h["pos"] + 160].strip()
            ini_nombre = h["pos"] - len(h["nombre"] or "") - 4
            antes = t[max(0, ini_nombre - 60):max(0, ini_nombre)]
            antes = re.split(r"[.;:]\s", antes)[-1]  # solo la misma frase
            personas.append({**h, "vinculo": _vinculo(antes, h["fecha_nacimiento"]), "contexto": contexto})
        else:
            for k in ("nombre", "dni", "fecha_nacimiento"):
                if not igual[k] and h[k]:
                    igual[k] = h[k]
    for p in personas:
        p.pop("pos", None)
        if p["fecha_nacimiento"]:
            p["fecha_nacimiento"] = p["fecha_nacimiento"].isoformat()
    return personas
